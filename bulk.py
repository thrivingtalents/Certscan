"""Reading a list of names (and optional emails) from an Excel or CSV file for bulk import."""
import csv
import io
import re

from openpyxl import load_workbook

MAX_ROWS = 2000
NAME_HEADERS = ("full name", "fullname", "name", "nama penuh", "nama", "participant", "peserta", "attendee")
EMAIL_HEADERS = ("email", "e-mail", "emel", "e-mel", "email address", "alamat emel")


class BulkImportError(ValueError):
    pass


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _read_rows(data, filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("xlsx", "xlsm"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise BulkImportError("Could not open that Excel file. Save it again as .xlsx and retry.") from exc
        ws = wb.active
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if ext == "xls":
        raise BulkImportError("Old .xls files aren't supported. In Excel choose File → Save As → .xlsx, then upload that.")
    if ext in ("csv", "txt"):
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [[_cell(v) for v in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise BulkImportError("Upload an Excel (.xlsx) or CSV file.")


def _find_columns(rows, email_re):
    """Return (header_row_index or None, name_col, email_col or None)."""
    for i, row in enumerate(rows[:5]):
        low = [c.lower() for c in row]
        name_col = next((j for h in NAME_HEADERS for j, c in enumerate(low) if c == h), None)
        if name_col is None:
            name_col = next((j for j, c in enumerate(low) if "name" in c or "nama" in c), None)
        email_col = next((j for h in EMAIL_HEADERS for j, c in enumerate(low) if c == h), None)
        if email_col is None:
            email_col = next((j for j, c in enumerate(low) if "mail" in c or "emel" in c), None)
        if name_col is not None and name_col != email_col:
            return i, name_col, email_col
    # no header row: names are in the first column with text, emails in the first column holding addresses
    width = max((len(r) for r in rows), default=0)
    email_col = next((j for j in range(width) if any(j < len(r) and email_re.match(r[j]) for r in rows)), None)
    name_col = next((j for j in range(width) if j != email_col and any(j < len(r) and r[j] for r in rows)), None)
    if name_col is None:
        raise BulkImportError("The file has no names in it.")
    return None, name_col, email_col


def read_people(data, filename, email_re):
    """Return (people, problems).

    people: list of dicts {row, name, email} where email may be ''.
    problems: list of (row, message) for rows that were skipped or changed.
    """
    rows = _read_rows(data, filename)
    if not any(any(r) for r in rows):
        raise BulkImportError("The file is empty.")
    header, name_col, email_col = _find_columns(rows, email_re)
    people, problems = [], []
    start = 0 if header is None else header + 1
    for i in range(start, len(rows)):
        r = rows[i]
        row_no = i + 1  # as numbered in Excel
        name = r[name_col] if name_col < len(r) else ""
        email = (r[email_col] if email_col is not None and email_col < len(r) else "").lower()
        if not name and not email:
            continue
        if len(people) >= MAX_ROWS:
            raise BulkImportError(f"That's more than {MAX_ROWS} names. Split the file and upload each part.")
        if email and not email_re.match(email):
            problems.append((row_no, f"“{email}” isn't a valid email, so the certificate won't be emailed"))
            email = ""
        people.append({"row": row_no, "name": name, "email": email})
    if not people:
        raise BulkImportError("No names found under the name column.")
    return people, problems


def example_csv():
    return "﻿Name,Email\nNur Aisyah binti Rahman,aisyah@example.com\nTan Wei Ming,\nAnand a/l Murugan,anand@example.com\n"
